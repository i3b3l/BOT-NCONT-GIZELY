# -*- coding: utf-8 -*-
# Por Izabel Galvão :3
"""
PROCEDIMENTO: CERTIFICAR E LIQUIDAR COM DATA DE PAGAMENTO
Versão revisada + UP de pré-validação, alinhamento das abas e segurança operacional.
Nesta versão, Data de Emissão, Data de Aceite e Data de Apresentação da CE
usam a Data do Pagamento informada na coluna 17 da aba Entrada.

Principais melhorias:
- Não usa delete_rows()+append() para atualizar CE/NL.
- Preserva os dados em uma única planilha de recuperação se a gravação falhar.
- Relê Data do Pagamento em cada linha da certificação e da liquidação.
- Usa a Data do Pagamento nas datas de Emissão, Aceite e Apresentação da CE.
- Corrige campos de data mascarados do SIGEF usando fill() e fallback direto com validação.
- Remove anos fixos de CE/NL e usa o exercício do documento.
- Valida CPF e nome na MESMA linha retornada pelo SIGEF.
- Não usa clipboard para capturar o número da CE.
- Valida datas, CPF, processo, empenho e valor antes de registrar.
- Evita substituir data inválida silenciosamente pela data atual.
- Inicializa estados por linha.
- Melhora tratamento de erros e fechamento seguro.
- Usa seletores/fluxo do código original sempre que possível.
- Pré-valida a planilha inteira antes de acessar o SIGEF.
- Valida cabeçalhos, tipos de texto, duplicidades e alinhamento entre abas.
- Mantém CE/NL na mesma linha e replica os dados-base sem delete_rows/append.

ATENÇÃO:
Abra o Chrome com depuração remota antes de executar:
"C:\\Program Files\\Google\\Chrome\\Application\\chrome.exe" --remote-debugging-port=9222 --user-data-dir="C:\\ChromeDebugProfile"

Depois, faça login no SIGEF.
"""

from __future__ import annotations

import re
import shutil
import sys
import time
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import keyboard
import openpyxl
import pyautogui
from playwright.sync_api import  (
    TimeoutError as PlaywrightTimeoutError,
    sync_playwright,
)

# ============================================================
# CONFIGURAÇÕES
# ============================================================

PLANILHA = str(Path(__file__).resolve().with_name("Pagamentos_Tratado.xlsx"))
BACKUP_FIXO = "Pagamentos_Backup.xlsx"
PASTA_RECUPERACAO = "RECUPERACAO"

CHROME_DEBUG_URL = "http://localhost:9222"

TECLA_DE_PANICO = "Esc"

STATUS_PESQUISAR_CE = "PESQUISAR_CE"
STATUS_PESQUISAR_NL = "PESQUISAR_NL"

COL_UG = 1
COL_GESTAO = 2
COL_PROCESSO = 3
COL_NOME = 4
COL_CPF = 5
COL_VALOR = 6
COL_BANCO = 7
COL_AGENCIA = 8
COL_CONTA = 9
COL_EMPENHO = 10
COL_CE = 11
COL_NL = 12
COL_PP = 13
COL_OB = 14
COL_DATA = 15
COL_OBSERVACAO = 16
COL_DATA_PAGAMENTO = 17

MESES = {
    "01": "Janeiro",
    "02": "Fevereiro",
    "03": "Março",
    "04": "Abril",
    "05": "Maio",
    "06": "Junho",
    "07": "Julho",
    "08": "Agosto",
    "09": "Setembro",
    "10": "Outubro",
    "11": "Novembro",
    "12": "Dezembro",
}

robo_deve_parar = False


# ============================================================
# EXCEÇÕES
# ============================================================

class ErroAutomacao(Exception):
    """Erro controlado da automação."""


class ErroDados(ErroAutomacao):
    """Erro nos dados da planilha."""


class ErroValidacaoSIGEF(ErroAutomacao):
    """Erro de validação de dados encontrados no SIGEF."""


# ============================================================
# TECLA DE PÂNICO / ENCERRAMENTO
# ============================================================

# Ative somente para investigar detalhes de preenchimento e validação.
MOSTRAR_DETALHES = False


def mensagem_detalhada(*args, **kwargs):
    if MOSTRAR_DETALHES:
        print(*args, **kwargs)


def parar_execucao():
    global robo_deve_parar
    print("\n!!! TECLA ESC ACIONADA! ENCERRANDO AUTOMAÇÃO COM SEGURANÇA !!!")
    robo_deve_parar = True


def verificar_panico():
    if robo_deve_parar:
        raise KeyboardInterrupt("Tecla ESC acionada.")


# ============================================================
# UTILITÁRIOS
# ============================================================

def caminho_backup_recuperacao():
    principal = Path(PLANILHA).resolve()
    return principal.parent / PASTA_RECUPERACAO / BACKUP_FIXO


def recuperar_backup_pendente():
    backup = caminho_backup_recuperacao()
    if not backup.exists():
        return
    resposta = pyautogui.confirm(
        text=("Existe uma planilha de recuperação criada após uma falha ao salvar.\n"
              "Feche a planilha no Excel. Deseja recuperar esses dados na planilha principal?\n"
              "Cancelar encerra o robô sem alterar os arquivos."),
        title="Recuperar planilha após erro",
        buttons=["Recuperar", "Cancelar"],
    )
    if resposta != "Recuperar":
        raise ErroDados("Recuperação pendente. Automação encerrada sem novas operações.")
    # Verificar que é um workbook legível antes de substituir a principal.
    teste = openpyxl.load_workbook(backup, read_only=True)
    teste.close()
    principal = Path(PLANILHA).resolve()
    temporario = principal.with_name(principal.stem + ".recuperacao.tmp.xlsx")
    shutil.copy2(backup, temporario)
    temporario.replace(principal)
    backup.unlink()
    print("[RECUPERADO] Dados da última falha restaurados na planilha principal.")


def salvar_planilha_segura(book):
    principal = Path(PLANILHA).resolve()
    backup = caminho_backup_recuperacao()
    temporario = principal.with_name(principal.stem + ".gravacao.tmp.xlsx")
    try:
        book.save(temporario)
        temporario.replace(principal)
    except Exception as erro_principal:
        # Salvar o estado atual, incluindo o número recém-gerado, não uma
        # cópia antiga do arquivo que falhou. Sempre o mesmo nome de backup.
        backup.parent.mkdir(parents=True, exist_ok=True)
        temporario_backup = backup.with_name(backup.stem + ".tmp.xlsx")
        try:
            book.save(temporario_backup)
            temporario_backup.replace(backup)
        except Exception as erro_backup:
            raise ErroDados(
                f"Falha na principal ({erro_principal}) e no backup ({erro_backup}). "
                "Confira o último documento no SIGEF antes de repetir a operação."
            ) from erro_backup
        finally:
            # Só limpar a tentativa primária se a recuperação foi preservada.
            if backup.exists() and temporario.exists():
                try:
                    temporario.unlink()
                except OSError:
                    pass
        raise ErroDados(
            f"Não foi possível salvar {principal.name}: {erro_principal}. "
            f"Os dados atuais foram preservados em {backup.name}. "
            "Feche o Excel; na próxima execução escolha Recuperar."
        ) from erro_principal


def texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def limpar_observacao(valor) -> str:
    observacao = texto(valor)

    # Remove caracteres invisíveis que podem vir de textos copiados da web
    # e que o SIGEF exibe como entidades, por exemplo: &#8203;.
    observacao = re.sub(
        r"&#(?:8203|8204|8205|8288|65279);",
        "",
        observacao,
        flags=re.IGNORECASE,
    )
    observacao = re.sub(r"[\u200B\u200C\u200D\u2060\uFEFF]", "", observacao)

    # Converte espaços não separáveis em espaços comuns sem alterar o texto.
    observacao = observacao.replace("\u00A0", " ")
    return observacao.strip()


def somente_digitos(valor) -> str:
    return re.sub(r"\D", "", texto(valor))


def normalizar_nome(valor: str) -> str:
    valor = texto(valor).upper()
    valor = unicodedata.normalize("NFKD", valor)
    valor = "".join(c for c in valor if not unicodedata.combining(c))
    valor = re.sub(r"\s+", " ", valor).strip()
    return valor


def primeiro_nome_normalizado(nome: str) -> str:
    nome_limpo = normalizar_nome(nome)
    if not nome_limpo:
        return ""
    return nome_limpo.split()[0]


def formatar_cpf(valor) -> tuple[str, str]:
    digitos = somente_digitos(valor)

    if not digitos:
        raise ErroDados("CPF não preenchido.")

    digitos = digitos.zfill(11)

    if len(digitos) != 11:
        raise ErroDados(f"CPF inválido: {valor!r}")

    formatado = f"{digitos[:3]}.{digitos[3:6]}.{digitos[6:9]}-{digitos[9:]}"
    return digitos, formatado


def formatar_processo(valor) -> str:
    digitos = somente_digitos(valor)

    if not digitos:
        raise ErroDados("Número do processo não preenchido.")

    digitos = digitos.zfill(16)

    if len(digitos) != 16:
        raise ErroDados(
            f"Processo inválido: {valor!r}. Esperados 16 dígitos."
        )

    return f"{digitos[:4]}.{digitos[4:10]}/{digitos[10:14]}-{digitos[14:]}"


def formatar_data(valor, nome_campo: str) -> str:
    if valor is None or texto(valor) == "":
        raise ErroDados(f"{nome_campo} não preenchida.")

    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")

    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")

    bruto = texto(valor)

    formatos = (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d.%m.%Y",
        "%Y-%m-%d",
        "%Y/%m/%d",
    )

    candidato = bruto[:10]

    for formato in formatos:
        try:
            return datetime.strptime(candidato, formato).strftime("%d/%m/%Y")
        except ValueError:
            pass

    digitos = somente_digitos(bruto)

    if len(digitos) == 8:
        try:
            return datetime.strptime(digitos, "%d%m%Y").strftime("%d/%m/%Y")
        except ValueError:
            pass

    raise ErroDados(
        f"{nome_campo} inválida: {bruto!r}. "
        "Corrija a planilha antes de continuar."
    )



def formatar_empenho(valor) -> tuple[int, int, str]:
    bruto = texto(valor).upper().replace(" ", "")

    if not bruto:
        raise ErroDados("Nota de Empenho não preenchida.")

    match = re.fullmatch(r"(\d{4})NE(\d{1,6})", bruto)

    if not match:
        raise ErroDados(
            f"Nota de Empenho inválida: {valor!r}. "
            "Informe no formato AAAANE000000, por exemplo 2026NE000123."
        )

    exercicio = int(match.group(1))
    numero = int(match.group(2))
    nota_formatada = f"{exercicio}NE{numero:06d}"

    return exercicio, numero, nota_formatada

def decimal_monetario(valor) -> Decimal:
    bruto = texto(valor)

    if not bruto:
        raise ErroDados("Valor não preenchido.")

    bruto = bruto.replace("R$", "").replace(" ", "")

    if "," in bruto and "." in bruto:
        # Ex.: 1.234,56
        if bruto.rfind(",") > bruto.rfind("."):
            bruto = bruto.replace(".", "").replace(",", ".")
        # Ex.: 1,234.56
        else:
            bruto = bruto.replace(",", "")
    elif "," in bruto:
        bruto = bruto.replace(".", "").replace(",", ".")

    try:
        return Decimal(bruto)
    except InvalidOperation as exc:
        raise ErroDados(f"Valor monetário inválido: {valor!r}") from exc


def valor_para_sigef(valor) -> str:
    """
    Mantém o comportamento próximo ao original.
    Se o Excel trouxer número, evita notação científica.
    """
    if valor is None:
        raise ErroDados("Valor não preenchido.")

    if isinstance(valor, Decimal):
        return format(valor, "f")

    if isinstance(valor, int):
        return str(valor)

    if isinstance(valor, float):
        return format(Decimal(str(valor)), "f")

    return texto(valor)


def prefixo_ce(exercicio: int) -> str:
    return f"{exercicio}CE"


def extrair_numero_ce(ce: str) -> str:
    bruto = texto(ce).upper()

    match = re.fullmatch(r"(\d{4})CE(\d+)", bruto)
    if not match:
        raise ErroDados(f"CE inválida: {ce!r}")

    return match.group(2)


def extrair_exercicio_ce(ce: str) -> int:
    bruto = texto(ce).upper()
    match = re.fullmatch(r"(\d{4})CE(\d+)", bruto)

    if not match:
        raise ErroDados(f"CE inválida: {ce!r}")

    return int(match.group(1))


def aguardar_rede(pagina, timeout=30000):
    try:
        pagina.wait_for_load_state("networkidle", timeout=timeout)
    except PlaywrightTimeoutError:
        print(
            "[AVISO] A página não atingiu 'networkidle' no prazo. "
            "Continuando porque o elemento seguinte será validado."
        )


def preencher_sequencial(locator, valor: str, limpar=True):
    verificar_panico()
    locator.wait_for(state="visible", timeout=10000)

    if limpar:
        try:
            locator.press("Control+A")
            locator.press("Backspace")
        except Exception:
            pass

    locator.press_sequentially(str(valor))


def preencher_data_sigef(locator, data_formatada: str):
    """
    Preenche campos de data mascarados do SIGEF com segurança.

    Estratégia:
    1. Valida DD/MM/AAAA.
    2. Tenta preencher diretamente com Playwright fill(), evitando que
       a máscara processe cada tecla individualmente.
    3. Valida o valor retornado pelo próprio campo.
    4. Se a máscara ainda alterar a data, aplica o valor diretamente no
       elemento HTML e dispara eventos input/change/blur.
    5. Só continua se o campo terminar exatamente com a data esperada.
    """
    verificar_panico()

    try:
        data_obj = datetime.strptime(data_formatada, "%d/%m/%Y")
    except ValueError as exc:
        raise ErroDados(
            f"Data inválida para preenchimento no SIGEF: {data_formatada!r}"
        ) from exc

    esperado_formatado = data_obj.strftime("%d/%m/%Y")
    esperado_numeros = data_obj.strftime("%d%m%Y")

    locator.wait_for(state="visible", timeout=10000)

    def data_esta_correta():
        try:
            valor_atual = texto(locator.input_value())
        except Exception:
            return False, ""

        return (
            somente_digitos(valor_atual) == esperado_numeros,
            valor_atual,
        )

    # --------------------------------------------------------
    # TENTATIVA 1: fill() com a data completa DD/MM/AAAA
    # --------------------------------------------------------
    try:
        locator.click()
        locator.fill(esperado_formatado)
        locator.press("Tab")
        time.sleep(0.4)

        correta, valor_final = data_esta_correta()

        if correta:
            mensagem_detalhada(
                f"[DATA OK] {esperado_formatado} preenchida com fill() "
                f"(campo retornou {valor_final!r})."
            )
            return

        print(
            f"[DATA] fill() foi alterado pela máscara do SIGEF: "
            f"{valor_final!r}. Tentando preenchimento direto..."
        )

    except Exception as erro_fill:
        print(
            f"[DATA] fill() não funcionou neste campo: {erro_fill}. "
            "Tentando preenchimento direto..."
        )

    # --------------------------------------------------------
    # TENTATIVA 2: define o valor no elemento HTML e dispara
    # eventos que o SIGEF costuma usar para validar datas.
    # --------------------------------------------------------
    try:
        locator.evaluate(
            """
            (el, valor) => {
                el.focus();
                el.value = valor;

                el.dispatchEvent(new Event('input', {
                    bubbles: true
                }));

                el.dispatchEvent(new Event('change', {
                    bubbles: true
                }));

                el.dispatchEvent(new Event('blur', {
                    bubbles: true
                }));

                el.blur();
            }
            """,
            esperado_formatado,
        )

        time.sleep(0.5)

        correta, valor_final = data_esta_correta()

        if correta:
            mensagem_detalhada(
                f"[DATA OK] {esperado_formatado} preenchida diretamente "
                f"(campo retornou {valor_final!r})."
            )
            return

    except Exception as erro_js:
        valor_final = ""
        print(
            f"[DATA] O preenchimento direto também apresentou erro: {erro_js}"
        )

    # --------------------------------------------------------
    # PROTEÇÃO FINAL
    # --------------------------------------------------------
    raise ErroValidacaoSIGEF(
        "O SIGEF não recebeu a data corretamente após duas tentativas. "
        f"Esperado: {esperado_formatado} | "
        f"Campo ficou: {valor_final!r}. "
        "A operação foi interrompida antes da inclusão."
    )


def obter_valor_locator(locator) -> str:
    try:
        return texto(locator.input_value())
    except Exception:
        return texto(locator.inner_text())


def mensagem_erro_linha(linha: int, erro: Exception) -> str:
    return f"Erro na linha {linha}: {erro}"


# ============================================================
# PLANILHA
# ============================================================

def abrir_planilha():
    recuperar_backup_pendente()
    book = openpyxl.load_workbook(PLANILHA)

    abas_obrigatorias = [
        "Entrada",
        "Despesas Certificadas",
        "Notas de Liquidação",
        "Preparações de Pagamento",
        "Ordens Bancárias",
        "Saída",
    ]

    faltando = [aba for aba in abas_obrigatorias if aba not in book.sheetnames]

    if faltando:
        book.close()
        raise ErroDados(
            "Abas obrigatórias ausentes na planilha: "
            + ", ".join(faltando)
        )

    return {
        "book": book,
        "entrada": book["Entrada"],
        "ce": book["Despesas Certificadas"],
        "nl": book["Notas de Liquidação"],
        "pp": book["Preparações de Pagamento"],
        "ob": book["Ordens Bancárias"],
        "saida": book["Saída"],
    }



def linha_vazia(entrada, linha: int) -> bool:
    return all(
        entrada.cell(linha, coluna).value is None
        or texto(entrada.cell(linha, coluna).value) == ""
        for coluna in range(1, COL_DATA_PAGAMENTO + 1)
    )


def ultima_linha_com_dados(entrada) -> int:
    """
    Retorna a última linha realmente preenchida da aba Entrada.
    Não confia apenas em max_row, pois formatações do Excel podem inflar esse valor.
    """
    for linha in range(entrada.max_row, 1, -1):
        if not linha_vazia(entrada, linha):
            return linha
    return 1


def normalizar_cabecalho(valor) -> str:
    valor = normalizar_nome(texto(valor))
    return re.sub(r"\s+", " ", valor).strip()


def validar_cabecalhos(planilhas):
    esperados_entrada = {
        1: "UNIDADE GESTORA",
        2: "GESTAO",
        3: "PROCESSO Nº",
        4: "NOME",
        5: "CPF",
        6: "VALOR",
        7: "BANCO",
        8: "AGENCIA",
        9: "CONTA CORRENTE",
        10: "EMPENHO",
        11: "DESPESA CERTIFICADA",
        12: "NOTA DE LIQUIDAÇÃO",
        13: "PREPARAÇÃO DE PAGAMENTO",
        14: "ORDEM BANCÁRIA",
        15: "DATA",
        16: "OBSERVAÇÃO",
        17: "DATA DO PAGAMENTO",
    }

    esperados_saida = {k: v for k, v in esperados_entrada.items() if k <= 16}

    erros = []

    for coluna, esperado in esperados_entrada.items():
        encontrado = planilhas["entrada"].cell(1, coluna).value
        if normalizar_cabecalho(encontrado) != normalizar_cabecalho(esperado):
            erros.append(
                f"Entrada!{planilhas['entrada'].cell(1, coluna).coordinate}: "
                f"esperado {esperado!r}, encontrado {encontrado!r}."
            )

    for chave in ("ce", "nl", "pp", "ob", "saida"):
        ws = planilhas[chave]
        for coluna, esperado in esperados_saida.items():
            encontrado = ws.cell(1, coluna).value
            if normalizar_cabecalho(encontrado) != normalizar_cabecalho(esperado):
                erros.append(
                    f"{ws.title}!{ws.cell(1, coluna).coordinate}: "
                    f"esperado {esperado!r}, encontrado {encontrado!r}."
                )

    return erros


def validar_identificadores_como_texto(entrada, linha: int):
    """
    Campos identificadores devem estar armazenados como texto no Excel.
    Isso protege zeros à esquerda e evita alterações silenciosas.
    """
    campos = {
        COL_UG: "UNIDADE GESTORA",
        COL_GESTAO: "GESTAO",
        COL_PROCESSO: "PROCESSO Nº",
        COL_CPF: "CPF",
        COL_VALOR: "VALOR",
        COL_BANCO: "BANCO",
        COL_AGENCIA: "AGENCIA",
        COL_CONTA: "CONTA CORRENTE",
        COL_EMPENHO: "EMPENHO",
    }

    erros = []

    for coluna, nome in campos.items():
        valor = entrada.cell(linha, coluna).value
        if valor is None or texto(valor) == "":
            continue

        if not isinstance(valor, str):
            erros.append(
                f"Linha {linha} - {nome}: deve estar armazenado como TEXTO "
                f"no Excel; tipo encontrado: {type(valor).__name__}."
            )

    return erros


def chave_identidade_registro(registro: dict):
    return (
        registro["processo_formatado"],
        registro["cpf_formatado"],
        registro["nota_empenho"],
        registro["valor_decimal"],
    )


def validar_alinhamento_saida(planilhas, registro: dict):
    """
    Se uma aba de saída já possuir dados na mesma linha, confirma que eles
    pertencem ao mesmo registro da aba Entrada.
    """
    linha = registro["linha"]
    erros = []

    for chave in ("ce", "nl", "pp", "ob"):
        ws = planilhas[chave]

        possui_dados = any(
            ws.cell(linha, coluna).value is not None
            and texto(ws.cell(linha, coluna).value) != ""
            for coluna in range(1, 15)
        )

        if not possui_dados:
            continue

        processo_saida = somente_digitos(ws.cell(linha, COL_PROCESSO).value)
        cpf_saida = somente_digitos(ws.cell(linha, COL_CPF).value)

        processo_entrada = somente_digitos(registro["processo_formatado"])
        cpf_entrada = somente_digitos(registro["cpf_formatado"])

        if processo_saida and processo_saida != processo_entrada:
            erros.append(
                f"Linha {linha} - aba '{ws.title}' desalinhada: "
                "PROCESSO diferente da aba Entrada."
            )

        if cpf_saida and cpf_saida.zfill(11) != cpf_entrada.zfill(11):
            erros.append(
                f"Linha {linha} - aba '{ws.title}' desalinhada: "
                "CPF diferente da aba Entrada."
            )

    return erros


def validar_documento_existente(valor, tipo: str, linha: int):
    bruto = texto(valor).upper()

    if not bruto:
        return None

    if tipo == "CE":
        if bruto == STATUS_PESQUISAR_CE:
            return None
        if not re.fullmatch(r"\d{4}CE\d{6}", bruto):
            return f"Linha {linha} - CE em formato inválido: {valor!r}."

    if tipo == "NL":
        if bruto == STATUS_PESQUISAR_NL:
            return None
        if not re.fullmatch(r"\d{4}NL\d{6}", bruto):
            return f"Linha {linha} - NL em formato inválido: {valor!r}."

    return None


def prevalidar_planilha(planilhas) -> int:
    print("\n" + "=" * 60)
    print("PRÉ-VALIDAÇÃO DA PLANILHA")
    print("=" * 60)

    erros = []
    avisos = []

    erros.extend(validar_cabecalhos(planilhas))

    ultima_linha = ultima_linha_com_dados(planilhas["entrada"])

    if ultima_linha < 2:
        raise ErroDados("A aba Entrada não possui registros para processar.")

    total_registros = 0
    ces_vistas = {}
    nls_vistas = {}
    identidades = {}

    for linha in range(2, ultima_linha + 1):
        if linha_vazia(planilhas["entrada"], linha):
            avisos.append(f"Linha {linha} está vazia e será ignorada.")
            continue

        total_registros += 1
        erros.extend(validar_identificadores_como_texto(planilhas["entrada"], linha))

        try:
            registro = ler_registro(planilhas, linha)
            registro = validar_e_formatar_registro(registro)
        except Exception as erro:
            erros.append(f"Linha {linha}: {erro}")
            continue

        erros.extend(validar_alinhamento_saida(planilhas, registro))

        erro_ce = validar_documento_existente(
            registro["despesa_certificada"], "CE", linha
        )
        if erro_ce:
            erros.append(erro_ce)

        erro_nl = validar_documento_existente(
            registro["liquidacao"], "NL", linha
        )
        if erro_nl:
            erros.append(erro_nl)

        ce = texto(registro["despesa_certificada"]).upper()
        if ce and ce != STATUS_PESQUISAR_CE:
            if ce in ces_vistas:
                erros.append(
                    f"CE duplicada: {ce} aparece nas linhas "
                    f"{ces_vistas[ce]} e {linha}."
                )
            else:
                ces_vistas[ce] = linha

        nl = texto(registro["liquidacao"]).upper()
        if nl and nl != STATUS_PESQUISAR_NL:
            if nl in nls_vistas:
                erros.append(
                    f"NL duplicada: {nl} aparece nas linhas "
                    f"{nls_vistas[nl]} e {linha}."
                )
            else:
                nls_vistas[nl] = linha

        identidade = chave_identidade_registro(registro)
        if identidade in identidades:
            avisos.append(
                f"Possível registro duplicado nas linhas "
                f"{identidades[identidade]} e {linha} "
                "(mesmo processo, CPF, empenho e valor)."
            )
        else:
            identidades[identidade] = linha

    # Verifica se alguma aba possui conteúdo além da última linha útil da Entrada.
    for chave in ("ce", "nl", "pp", "ob"):
        ws = planilhas[chave]
        for linha in range(ultima_linha + 1, ws.max_row + 1):
            if any(
                ws.cell(linha, coluna).value is not None
                and texto(ws.cell(linha, coluna).value) != ""
                for coluna in range(1, 15)
            ):
                avisos.append(
                    f"A aba '{ws.title}' possui dados na linha {linha}, "
                    "além da última linha útil da Entrada."
                )
                break

    print(f"Última linha útil da Entrada: {ultima_linha}")
    print(f"Registros encontrados: {total_registros}")
    print(f"CEs existentes: {len(ces_vistas)}")
    print(f"NLs existentes: {len(nls_vistas)}")

    if avisos:
        print("\n[AVISOS]")
        for aviso in avisos:
            print(f" - {aviso}")

    if erros:
        print("\n[ERROS]")
        for erro in erros:
            print(f" - {erro}")

        raise ErroDados(
            f"A pré-validação encontrou {len(erros)} erro(s). "
            "Corrija a planilha antes de acessar o SIGEF."
        )

    print("\nCabeçalhos: OK")
    print("Tipos de texto: OK")
    print("Processos/CPFs/Empenhos/Valores/Datas: OK")
    print("Duplicidade de CE/NL: OK")
    print("Alinhamento entre abas: OK")
    print("\nPLANILHA APROVADA. Iniciando SIGEF...")
    print("=" * 60 + "\n")

    return ultima_linha


def ler_registro(planilhas, linha: int) -> dict:
    entrada = planilhas["entrada"]

    registro = {
        "linha": linha,
        "ug": texto(entrada.cell(linha, COL_UG).value),
        "gestao": texto(entrada.cell(linha, COL_GESTAO).value),
        "processo_raw": entrada.cell(linha, COL_PROCESSO).value,
        "nome": texto(entrada.cell(linha, COL_NOME).value),
        "cpf_raw": entrada.cell(linha, COL_CPF).value,
        "valor_raw": entrada.cell(linha, COL_VALOR).value,
        "banco": texto(entrada.cell(linha, COL_BANCO).value),
        "agencia": texto(entrada.cell(linha, COL_AGENCIA).value),
        "conta": texto(entrada.cell(linha, COL_CONTA).value),
        "empenho_raw": entrada.cell(linha, COL_EMPENHO).value,
        "despesa_certificada": texto(
            planilhas["ce"].cell(linha, COL_CE).value
        ),
        "liquidacao": texto(
            planilhas["nl"].cell(linha, COL_NL).value
        ),
        "preparacao_pagamento": texto(
            planilhas["pp"].cell(linha, COL_PP).value
        ),
        "ordem_bancaria": texto(
            planilhas["ob"].cell(linha, COL_OB).value
        ),
        "data_raw": entrada.cell(linha, COL_DATA).value,
        "observacao": limpar_observacao(
            entrada.cell(linha, COL_OBSERVACAO).value
        ),
        "data_pagamento_raw": entrada.cell(
            linha, COL_DATA_PAGAMENTO
        ).value,
    }

    return registro


def validar_e_formatar_registro(registro: dict) -> dict:
    obrigatorios = {
        "UG": registro["ug"],
        "Gestão": registro["gestao"],
        "Nome": registro["nome"],
        "Observação": registro["observacao"],
    }

    vazios = [nome for nome, valor in obrigatorios.items() if not valor]

    if vazios:
        raise ErroDados(
            "Campos obrigatórios vazios: " + ", ".join(vazios)
        )

    cpf_sem_pontuacao, cpf_formatado = formatar_cpf(registro["cpf_raw"])
    processo_formatado = formatar_processo(registro["processo_raw"])
    data_pagamento = formatar_data(
        registro["data_pagamento_raw"],
        "Data do Pagamento",
    )
    # A Data do Pagamento é a única fonte de data do robô. A coluna DATA
    # permanece na planilha apenas para preservar a posição das demais colunas.
    data_formatada = data_pagamento
    exercicio, empenho_numero, nota_empenho = formatar_empenho(
        registro["empenho_raw"]
    )

    # Só valida conversão; o texto digitado permanece próximo ao original.
    valor_decimal = decimal_monetario(registro["valor_raw"])
    valor_digitacao = valor_para_sigef(registro["valor_raw"])

    processo_cortado = processo_formatado.split("/")[0].split(".")[1]
    numero_documento = f"{processo_cortado}-{registro['linha'] - 1}"

    mes = data_formatada.split("/")[1]
    competencia = MESES.get(mes)

    if not competencia:
        raise ErroDados(
            f"Mês inválido na Data do Pagamento da linha "
            f"{registro['linha']}: {data_pagamento}"
        )

    registro.update(
        {
            "cpf_sem_pontuacao": cpf_sem_pontuacao,
            "cpf_formatado": cpf_formatado,
            "processo_formatado": processo_formatado,
            "data_formatada": data_formatada,
            "data_pagamento": data_pagamento,
            "exercicio": exercicio,
            "empenho_numero": empenho_numero,
            "nota_empenho": nota_empenho,
            "valor_decimal": valor_decimal,
            "valor_digitacao": valor_digitacao,
            "numero_documento": numero_documento,
            "competencia": competencia,
            "primeiro_nome": primeiro_nome_normalizado(registro["nome"]),
        }
    )

    return registro


def montar_observacao_sigef(procedimento: str, registro: dict) -> str:
    observacao = registro["observacao"].rstrip()
    if not observacao.endswith((".", "!", "?")):
        observacao += "."

    return (
        f"{procedimento}: Pagamento para o(a) servidor(a) "
        f"{registro['cpf_formatado']} {registro['nome']} "
        f"referente à {observacao} "
        f"Processo Administrativo n: {registro['processo_formatado']}."
    )


def exibir_dados_registro(registro: dict, etapa: str):
    print(
        f"  CPF: {registro['cpf_formatado']} | "
        f"Processo: {registro['processo_formatado']}"
    )
    print(
        f"  Empenho: {registro['nota_empenho']} | "
        f"Valor: R$ {registro['valor_digitacao']}"
    )
    ano_competencia = registro["data_formatada"].split("/")[-1]
    print(
        f"  Competência: {registro['competencia']}/{ano_competencia} | "
        f"Data do pagamento: {registro['data_pagamento']}"
    )

    if etapa == "liquidacao":
        print(f"  Documento de origem: CE {registro['despesa_certificada']}")
    elif etapa == "preparacao":
        print(
            f"  Documentos: CE {registro['despesa_certificada']} | "
            f"NL {registro['liquidacao']}"
        )
        print(
            f"  Domicílio: banco {registro['banco']} | "
            f"agência {registro['agencia']} | conta {registro['conta']}"
        )



def _escrever_texto(ws, linha: int, coluna: int, valor):
    celula = ws.cell(linha, coluna)
    celula.value = "" if valor is None else str(valor)
    celula.number_format = "@"


def sincronizar_dados_base_saida(ws, registro: dict):
    valores = {
        COL_UG: registro["ug"],
        COL_GESTAO: registro["gestao"],
        COL_PROCESSO: registro["processo_formatado"],
        COL_NOME: registro["nome"],
        COL_CPF: registro["cpf_formatado"],
        COL_VALOR: registro["valor_digitacao"],
        COL_BANCO: registro["banco"],
        COL_AGENCIA: registro["agencia"],
        COL_CONTA: registro["conta"],
        COL_EMPENHO: registro["nota_empenho"],
    }

    for coluna, valor in valores.items():
        _escrever_texto(ws, registro["linha"], coluna, valor)


def salvar_ce(planilhas, registro: dict, ce: str):
    ws = planilhas["ce"]
    sincronizar_dados_base_saida(ws, registro)
    _escrever_texto(ws, registro["linha"], COL_CE, ce)
    salvar_planilha_segura(planilhas["book"])


def salvar_nl(planilhas, registro: dict, nl: str):
    ws = planilhas["nl"]
    sincronizar_dados_base_saida(ws, registro)
    _escrever_texto(
        ws,
        registro["linha"],
        COL_CE,
        registro["despesa_certificada"],
    )
    _escrever_texto(ws, registro["linha"], COL_NL, nl)
    salvar_planilha_segura(planilhas["book"])


def encontrar_guia_sigef(contexto):
    for pagina in contexto.pages:
        try:
            if "SIGEF" in pagina.url.upper():
                return pagina
        except Exception:
            pass

    raise ErroAutomacao(
        "Não encontrei nenhuma guia do SIGEF no Chrome. "
        "Abra o SIGEF, faça login e tente novamente."
    )

def abrir_funcionalidade(frame, guia, nome: str):
    verificar_panico()

    pesquisar = frame.get_by_placeholder(
        "Pesquisar funcionalidades do sistema..."
    )
    pesquisar.wait_for(state="visible", timeout=15000)
    pesquisar.press("Control+A")
    pesquisar.press("Backspace")
    pesquisar.press_sequentially(nome)

    funcionalidade = frame.get_by_title(nome)
    funcionalidade.wait_for(state="visible", timeout=15000)

    with guia.expect_popup(timeout=15000) as popup_info:
        funcionalidade.click()

    popup = popup_info.value
    aguardar_rede(popup)

    return popup


def localizar_e_validar_credor(pagina_favorecido, registro: dict):
    verificar_panico()

    pagina_favorecido.wait_for_load_state("domcontentloaded", timeout=30000)
    aguardar_rede(pagina_favorecido)

    print(
        f"  [CREDOR] Pesquisando pelo CPF: {registro['cpf_formatado']}",
        flush=True,
    )

    botao_cpf = pagina_favorecido.locator("#btnCPF")
    botao_cpf.wait_for(state="visible", timeout=10000)
    botao_cpf.click()

    preencher_cpf = pagina_favorecido.get_by_role("textbox").first
    preencher_sequencial(
        preencher_cpf,
        registro["cpf_sem_pontuacao"],
        limpar=True,
    )

    botao_confirmar = pagina_favorecido.get_by_role(
        "button",
        name="Confirmar a Consulta",
    )
    botao_confirmar.click()

    aguardar_rede(pagina_favorecido)

    codigo = pagina_favorecido.get_by_role(
        "cell",
        name=registro["cpf_formatado"],
        exact=True,
    )

    try:
        codigo.wait_for(state="visible", timeout=10000)
    except PlaywrightTimeoutError as exc:
        raise ErroValidacaoSIGEF(
            f"CPF {registro['cpf_formatado']} não encontrado no SIGEF."
        ) from exc

    # IMPORTANTE:
    # usa a própria célula do CPF encontrado como âncora,
    # em vez de pegar o primeiro CPF genérico da tabela.
    linha_correta = codigo.locator("xpath=..")
    celula_nome = linha_correta.locator("td").nth(1)

    nome_tela = texto(celula_nome.inner_text())
    primeiro_nome_tela = primeiro_nome_normalizado(nome_tela)
    primeiro_nome_esperado = registro["primeiro_nome"]

    mensagem_detalhada(f"Primeiro nome esperado: {primeiro_nome_esperado}")
    mensagem_detalhada(f"Primeiro nome encontrado na tela: {primeiro_nome_tela}")

    if primeiro_nome_tela != primeiro_nome_esperado:
        raise ErroValidacaoSIGEF(
            "Credor divergente. "
            f"Esperado: {primeiro_nome_esperado}; "
            f"encontrado: {primeiro_nome_tela or '(vazio)'}."
        )

    codigo.click()
    print(f"  [CREDOR] CPF e nome conferidos: {nome_tela}", flush=True)


# ============================================================
# CERTIFICAÇÃO
# ============================================================

def identificar_ce_gerada(pagina_ce):
    """Copia a CE completa exibida; nunca acrescenta ano a um sequencial."""
    padrao = re.compile(r"\b\d{4}CE\d{6}\b")
    limite = time.monotonic() + 10
    while time.monotonic() < limite:
        verificar_panico()
        numeros = set()
        # O resultado pode aparecer em um campo ou na mensagem de sucesso.
        campos = pagina_ce.locator("input:visible")
        for i in range(campos.count()):
            valor = campos.nth(i).input_value().strip()
            if padrao.fullmatch(valor):
                numeros.add(valor)
        # O SIGEF também exibe o identificador em dois campos adjacentes:
        # por exemplo, "2026CE" e "009188". Copiar ambos, sem usar ano da NE.
        sequenciais = pagina_ce.locator("#txtNuSeq:visible")
        if sequenciais.count() == 1:
            sequencial = sequenciais.nth(0)
            prefixos = sequencial.locator(
                "xpath=preceding::input[not(@type) or "
                "translate(@type, 'TEXT', 'text')='text'][1]"
            )
            if prefixos.count() == 1 and prefixos.nth(0).is_visible():
                prefixo = prefixos.nth(0).input_value().strip()
                numero = sequencial.input_value().strip()
                if re.fullmatch(r"\d{4}CE", prefixo) and re.fullmatch(r"\d{6}", numero):
                    numeros.add(prefixo + numero)
        mensagens = pagina_ce.get_by_text(padrao)
        for i in range(mensagens.count()):
            mensagem = mensagens.nth(i)
            if mensagem.is_visible():
                numeros.update(padrao.findall(mensagem.inner_text()))
        if len(numeros) == 1:
            return numeros.pop()
        if len(numeros) > 1:
            raise ErroValidacaoSIGEF(
                "Há mais de uma CE na tela; não é possível escolher automaticamente."
            )
        time.sleep(0.2)
    raise ErroValidacaoSIGEF(
        "O SIGEF não exibiu uma CE completa no formato AAAACE000000. "
        "O robô não irá inventar ou acrescentar o exercício."
    )


def certificar_linha(pagina_ce, planilhas, registro: dict, ultima_ce: str | None):
    verificar_panico()

    texto_da_ce = montar_observacao_sigef(
        "Certificação de Despesa",
        registro,
    )

    # REGRA DESTA VERSÃO:
    # Data de Emissão, Data de Aceite e Data de Apresentação da CE
    # recebem a Data do Pagamento da coluna 17 da aba Entrada.
    # A pré-validação já garante que a data exista e seja válida.
    data_emissao_final = registro["data_pagamento"]

    aguardar_rede(pagina_ce)

    print('  [1/5] Preenchendo documento e datas...', flush=True)
    campo_gestao = pagina_ce.locator("#txtCdGestao_SIGEFPesquisa")
    preencher_sequencial(campo_gestao, registro["gestao"])

    tipo_documento = pagina_ce.locator("#cmbCdTipoDocumento")
    tipo_documento.wait_for(state="visible", timeout=10000)
    tipo_documento.select_option(label="Outros")

    numero_documento = pagina_ce.locator("#txtNuDocumento")
    preencher_sequencial(
        numero_documento,
        registro["numero_documento"],
    )

    favorecido = pagina_ce.locator("#txtNmCredor_lnkBtnPesquisa")
    valor_documento = pagina_ce.locator("#txtVlDocumento")
    data_emissao = pagina_ce.locator("#txtDtEmissao_SIGEFData")
    data_aceite = pagina_ce.locator("#txtDtAceite_SIGEFData")
    data_apresentacao = pagina_ce.locator("#txtDtApresentacao_SIGEFData")
    competencia = pagina_ce.locator("#cboMesComp")
    observacao = pagina_ce.locator("#txtDeObservacao")
    atestado = pagina_ce.get_by_role(
        "checkbox",
        name="Sou responsável pelo atesto",
    )

    preencher_data_sigef(data_emissao, data_emissao_final)
    preencher_data_sigef(data_aceite, data_emissao_final)
    preencher_data_sigef(data_apresentacao, data_emissao_final)

    competencia.select_option(label=registro["competencia"])
    atestado.click()

    verificar_panico()

    print('  [2/5] Consultando e validando o credor...', flush=True)
    with pagina_ce.expect_popup(timeout=15000) as popup_info:
        favorecido.click()

    selecionar_favorecido = popup_info.value

    try:
        localizar_e_validar_credor(
            selecionar_favorecido,
            registro,
        )
    except Exception:
        try:
            if not selecionar_favorecido.is_closed():
                selecionar_favorecido.close()
        except Exception:
            pass
        raise

    aguardar_rede(pagina_ce)

    print('  [3/5] Preenchendo valor e observação...', flush=True)
    preencher_sequencial(
        valor_documento,
        registro["valor_digitacao"],
    )
    preencher_sequencial(observacao, texto_da_ce)

    botao_incluir = pagina_ce.get_by_role(
        "button",
        name="Incluir o Registro",
    )
    botao_incluir.wait_for(state="visible", timeout=10000)

    verificar_panico()
    print('  [4/5] Confirmando a certificação no SIGEF...', flush=True)
    botao_incluir.click()

    aguardar_rede(pagina_ce)

    documento_ja_cadastrado = False

    erro_na_tela = pagina_ce.get_by_role(
        "cell",
        name="Número Documento já cadastrado(a).",
        exact=True,
    )

    try:
        erro_na_tela.wait_for(state="visible", timeout=3000)
        documento_ja_cadastrado = True
    except PlaywrightTimeoutError:
        documento_ja_cadastrado = False

    if documento_ja_cadastrado:
        print(
            "[AVISO] Documento já cadastrado. "
            "A linha será marcada para pesquisa manual da CE."
        )
        salvar_ce(planilhas, registro, STATUS_PESQUISAR_CE,
        )

        botao_limpar = pagina_ce.get_by_role(
            "link",
            name="Limpar a Tela",
        )

        try:
            botao_limpar.click()
        except Exception:
            pass

        return STATUS_PESQUISAR_CE

    try:
        despesa_certificada = identificar_ce_gerada(pagina_ce)
    except Exception as exc:
        salvar_ce(planilhas, registro, STATUS_PESQUISAR_CE)
        raise ErroValidacaoSIGEF(
            "A CE pode ter sido gerada, mas seu número não foi identificado. "
            "Linha marcada PESQUISAR_CE para evitar nova certificação."
        ) from exc

    if ultima_ce and despesa_certificada == ultima_ce:
        raise ErroValidacaoSIGEF(
            f"A CE {despesa_certificada} foi repetida em duas linhas. "
            "O robô foi interrompido para evitar associação incorreta."
        )

    print('  [5/5] Gravando a CE gerada na planilha...', flush=True)
    salvar_ce(planilhas, registro, despesa_certificada,
    )

    print(
        f"  [SALVO] {despesa_certificada} "
    )

    botao_limpar = pagina_ce.get_by_role(
        "link",
        name="Limpar a Tela",
    )

    try:
        botao_limpar.click()
    except Exception:
        pass

    return despesa_certificada



def executar_certificacao(frame, guia, planilhas, ultima_linha: int):
    registros = [ler_registro(planilhas, linha)
                 for linha in range(2, ultima_linha + 1)
                 if not linha_vazia(planilhas["entrada"], linha)]
    existentes = sum(bool(r["despesa_certificada"]) and
                     r["despesa_certificada"] != STATUS_PESQUISAR_CE for r in registros)
    revisar = [r for r in registros if r["despesa_certificada"] == STATUS_PESQUISAR_CE]
    pendentes = [r for r in registros if not r["despesa_certificada"]]
    print("\n" + "=" * 64)
    print("CERTIFICAÇÃO DE DESPESAS")
    print(f"Já certificadas: {existentes} | A certificar: {len(pendentes)} ")
    print("=" * 64)
    for r in revisar:
        print(f"[ATENÇÃO] Linha {r['linha']} | {r['nome']} | Localizar CE no SIGEF.")
    if not pendentes:
        print("Nenhuma nova certificação a executar.")
        return
    pagina_ce = abrir_funcionalidade(frame, guia, "Manter Despesa Certificada")
    ultima_ce = None
    salvas = 0
    novas_revisoes = 0
    try:
        for indice, registro in enumerate(pendentes, start=1):
            verificar_panico()
            print(f"\n[{indice}/{len(pendentes)}] Linha {registro['linha']} | {registro['nome']}")
            registro = validar_e_formatar_registro(registro)
            exibir_dados_registro(registro, "certificacao")
            ultima_ce = certificar_linha(pagina_ce, planilhas, registro, ultima_ce)
            if ultima_ce == STATUS_PESQUISAR_CE:
                novas_revisoes += 1
            else:
                salvas += 1
    finally:
        try:
            pagina_ce.close()
        except Exception:
            pass
        print(f"\n[RESUMO] CEs salvas nesta etapa: {salvas} | Encaminhadas para conferência: {novas_revisoes}")


def preencher_valor_bruto_sigef(pagina, valor):
    """Mantém a sequência de teclado do robô original que funcionava."""
    esperado = decimal_monetario(valor)
    if not esperado.is_finite() or esperado <= 0:
        raise ErroDados("Valor bruto deve ser positivo e finito.")
    if esperado != esperado.quantize(Decimal("0.01")):
        raise ErroDados("Valor bruto deve ter no máximo duas casas decimais.")
    verificar_panico()
    campo = pagina.locator("#txtValorBrutoId")
    campo.wait_for(state="visible", timeout=10000)
    # Não limpar um campo vazio: preserva a inicialização da máscara original.
    if campo.input_value().strip():
        campo.press("Control+A")
        campo.press("Backspace")
    time.sleep(0.3)
    campo.press_sequentially(str(valor))
    time.sleep(0.3)
    # O próximo clique em Retenções tira o foco, como no código original.
    # A aceitação do valor será conferida depois da resposta do SIGEF.


def conferir_valor_bruto_apos_retencoes(pagina, esperado):
    campo = pagina.locator("#txtValorBrutoId")
    campo.wait_for(state="visible", timeout=10000)
    obtido = campo.input_value().strip()
    try:
        confere = decimal_monetario(obtido) == decimal_monetario(esperado)
    except ErroDados:
        confere = False
    if not confere:
        raise ErroValidacaoSIGEF(
            "O Valor Bruto foi apagado ou alterado após clicar em Retenções. "
            f"Esperado: {esperado}; campo após atualização: {obtido!r}. "
            "A consulta de retenções não será repetida automaticamente."
        )
    mensagem_detalhada(f"[VALOR BRUTO APÓS RETENÇÕES] {obtido}")


def esperar_valor_liquido(locator, timeout_segundos=10) -> str:
    inicio = time.monotonic()

    while time.monotonic() - inicio < timeout_segundos:
        verificar_panico()

        try:
            valor = texto(locator.input_value())
            if valor:
                return valor
        except Exception:
            pass

        time.sleep(0.2)

    raise ErroValidacaoSIGEF(
        "O SIGEF não calculou o valor líquido no tempo esperado."
    )


def localizar_nl_gerada(pagina_nl) -> str:
    # O exercício da NL pode ser diferente do exercício da CE e do empenho.
    # Capturar o número completo retornado na tela após a confirmação.
    padrao = re.compile(r"^\s*\d{4}NL\d{6}\s*$")
    celulas = pagina_nl.get_by_role("cell", name=padrao)
    celulas.first.wait_for(state="visible", timeout=30000)
    numeros = set()
    for indice in range(celulas.count()):
        celula = celulas.nth(indice)
        if celula.is_visible():
            numero = texto(celula.inner_text())
            if padrao.fullmatch(numero):
                numeros.add(numero.strip())
    if len(numeros) != 1:
        raise ErroValidacaoSIGEF(
            "Não foi possível identificar uma única NL na tela de confirmação. "
            f"Números encontrados: {sorted(numeros)}."
        )
    numero = numeros.pop()
    mensagem_detalhada(f"[NL CAPTURADA] {numero}")
    return numero


def preencher_ug_gestao_liquidacao(pagina_nl, ug: str, gestao: str):
    """Localiza a UG pelo campo de Gestão, sem depender do nome da secretaria."""
    verificar_panico()
    campo_gestao = pagina_nl.locator("#txtCdGestao_SIGEFPesquisa")
    campo_gestao.wait_for(state="visible", timeout=10000)
    # A UG é o campo de texto imediatamente anterior à Gestão.
    # Os controles podem ter tabelas internas distintas: não limitar a busca
    # ao ancestral tr da Gestão. Ignorar inputs ocultos e botões de pesquisa.
    campo_ug = campo_gestao.locator(
        "xpath=preceding::input[not(@type) or "
        "translate(@type, 'TEXT', 'text')='text'][1]"
    )
    campos = (("Unidade Gestora", campo_ug, str(ug)),
              ("Gestão", campo_gestao, str(gestao)))
    for nome, campo, esperado in campos:
        campo.wait_for(state="visible", timeout=10000)
        verificar_panico()
        # A UG pode vir preenchida e bloqueada pelo contexto do usuário.
        if campo.input_value().strip() != esperado:
            if not campo.is_enabled() or not campo.is_editable():
                raise ErroValidacaoSIGEF(
                    f"{nome} está bloqueada com valor diferente de {esperado}."
                )
            campo.fill(esperado)
            campo.press("Tab")
            aguardar_rede(pagina_nl)

    # A atualização da Gestão também pode alterar a UG: conferir ambas ao final.
    for nome, campo, esperado in campos:
        obtido = campo.input_value().strip()
        if obtido != esperado:
            raise ErroValidacaoSIGEF(
                f"{nome} incorreta na liquidação: esperado {esperado}, "
                f"encontrado {obtido!r}."
            )


def conferir_data_aceite_liquidacao(pagina, data_certificacao):
    campo = pagina.locator('input[id*="aceite" i]:visible')
    campo.wait_for(state="visible", timeout=10000)
    recebido = campo.input_value().strip()
    if recebido == data_certificacao:
        return
    if campo.is_enabled() and campo.is_editable():
        preencher_data_sigef(campo, data_certificacao)
        return
    raise ErroValidacaoSIGEF(
        f"A Data de Aceite da CE carregada pelo SIGEF é {recebido!r}, "
        f"mas a data da certificação na coluna 17 é {data_certificacao}. "
        "O campo está bloqueado; confira a CE antes de liquidar."
    )


def liquidar_linha(pagina_nl, planilhas, registro: dict):
    verificar_panico()

    ce = registro["despesa_certificada"]

    if ce == STATUS_PESQUISAR_CE:
        print(
            f"[PULAR] Linha {registro['linha']} precisa ter a CE "
            "localizada antes da liquidação."
        )
        return


    numero_ce = extrair_numero_ce(ce)

    texto_da_nl = montar_observacao_sigef(
        "Liquidação de Despesa",
        registro,
    )

    aguardar_rede(pagina_nl)

    print('  [1/6] Consultando a CE e preenchendo unidade e gestão...', flush=True)
    preencher_ug_gestao_liquidacao(
        pagina_nl, registro["ug"], registro["gestao"]
    )

    campo_ce = pagina_nl.locator(
        "#txtDespesaCertificadaNumero_SIGEFPesquisa"
    )
    preencher_sequencial(campo_ce, numero_ce)

    botao_pesquisar = pagina_nl.get_by_role(
        "button",
        name="Confirmar a Pesquisa",
    )
    botao_pesquisar.click()

    aguardar_rede(pagina_nl)

    verificar_panico()

    # A Data do Pagamento é relida e formatada PARA ESTA MESMA LINHA.
    print('  [2/6] Conferindo datas e adicionando o documento...', flush=True)
    data_vencimento = pagina_nl.locator(
        "#txtDataVencimento_SIGEFData"
    )
    preencher_data_sigef(
        data_vencimento,
        registro["data_pagamento"],
    )

    # A consulta da CE carrega a Data de Aceite nesta tela.
    # Se o SIGEF permitir edição, ajustar; se for somente leitura, conferir.
    conferir_data_aceite_liquidacao(pagina_nl, registro["data_pagamento"])

    adicionar = pagina_nl.get_by_role(
        "button",
        name="Adicionar Documento",
    )
    adicionar.click()

    aguardar_rede(pagina_nl)

    data_referencia = pagina_nl.locator(
        "#txtDtReferenciaId_SIGEFData"
    )
    data_referencia.wait_for(state="visible", timeout=10000)

    # Mesma origem das datas da certificação: coluna 17, Data do Pagamento.
    preencher_data_sigef(data_referencia, registro["data_pagamento"])

    ponto_interrogacao = pagina_nl.locator(
        "#txtNotaEmpenhoNumeroId_lnkBtnPesquisa"
    )

    print('  [3/6] Consultando e selecionando o empenho...', flush=True)
    with pagina_nl.expect_popup(timeout=15000) as popup_info:
        ponto_interrogacao.click()

    selecionar_empenho = popup_info.value

    try:
        aguardar_rede(selecionar_empenho)

        # O exercício usado na consulta vem da própria Nota de Empenho
        # informada na planilha, permitindo 2025, 2026 e anos futuros.
        exercicio_consulta = str(registro["exercicio"])

        # O ano é o campo de texto anterior ao número da NE no popup.
        # Seu identificador não precisa conter a palavra "exercicio".
        preencher_empenho = selecionar_empenho.locator(
            "#txtNotaEmpenhoNumero"
        )
        preencher_empenho.wait_for(state="visible", timeout=10000)
        preencher_exercicio = preencher_empenho.locator(
            "xpath=preceding::input[not(@type) or "
            "translate(@type, 'TEXT', 'text')='text'][1]"
        )
        preencher_exercicio.wait_for(state="visible", timeout=10000)
        preencher_exercicio.fill(exercicio_consulta)
        preencher_exercicio.press("Tab")
        aguardar_rede(selecionar_empenho)
        if preencher_exercicio.input_value().strip() != exercicio_consulta:
            raise ErroValidacaoSIGEF(
                "O SIGEF não manteve o exercício informado na consulta da NE."
            )
        preencher_sequencial(
            preencher_empenho,
            str(registro["empenho_numero"]),
        )

        botao_confirmar = selecionar_empenho.get_by_role(
            "button",
            name="Confirmar a Consulta",
        )
        botao_confirmar.click()

        aguardar_rede(selecionar_empenho)

        nota_empenho = selecionar_empenho.get_by_role(
            "cell",
            name=registro["nota_empenho"],
            exact=True,
        )

        try:
            nota_empenho.wait_for(
                state="visible",
                timeout=10000,
            )
        except PlaywrightTimeoutError as exc:
            raise ErroValidacaoSIGEF(
                f"Nota de Empenho {registro['nota_empenho']} "
                "não encontrada no SIGEF."
            ) from exc

        nota_empenho.click()

    except Exception:
        try:
            if not selecionar_empenho.is_closed():
                selecionar_empenho.close()
        except Exception:
            pass
        raise

    aguardar_rede(pagina_nl)

    print('  [4/6] Preenchendo o valor e conferindo retenções...', flush=True)
    preencher_valor_bruto_sigef(pagina_nl, registro["valor_digitacao"])

    botao_retencoes = pagina_nl.get_by_role(
        "button",
        name="Sugerir Retenções",
    )
    botao_retencoes.click()
    # Sincronização presente no robô original: aguardar a resposta da consulta.
    nao_existem_retencoes = pagina_nl.get_by_text("Não existem sugestões para")
    try:
        nao_existem_retencoes.wait_for(state="visible", timeout=30000)
    except PlaywrightTimeoutError as exc:
        raise ErroValidacaoSIGEF(
            "O SIGEF não exibiu a mensagem 'Não existem sugestões para' "
            "após consultar Retenções. Confira a resposta na tela; "
            "a liquidação não foi confirmada."
        ) from exc
    aguardar_rede(pagina_nl)

    valor_liquido = pagina_nl.locator("#txtValorLiquidoId")
    valor_liquido.wait_for(state="visible", timeout=10000)

    conferir_valor_bruto_apos_retencoes(pagina_nl, registro["valor_decimal"])
    try:
        valor_liquido_texto = esperar_valor_liquido(valor_liquido)
    except ErroValidacaoSIGEF as exc:
        conferir_valor_bruto_apos_retencoes(pagina_nl, registro["valor_decimal"])
        raise ErroValidacaoSIGEF(
            "O Valor Bruto permaneceu correto após Retenções, mas o SIGEF "
            "não retornou o Valor Líquido. Confira as mensagens e a aba "
            "Retenções na tela antes de continuar."
        ) from exc
    conferir_valor_bruto_apos_retencoes(pagina_nl, registro["valor_decimal"])
    valor_liquido_decimal = decimal_monetario(valor_liquido_texto)

    if valor_liquido_decimal != registro["valor_decimal"]:
        raise ErroValidacaoSIGEF(
            "Valor líquido diferente do valor da planilha. "
            f"Planilha: {registro['valor_decimal']} | "
            f"SIGEF: {valor_liquido_decimal}. "
            "A liquidação NÃO será confirmada automaticamente."
        )

    historico = pagina_nl.locator("#txtHistorico")
    preencher_sequencial(historico, texto_da_nl)

    botao_confirmacao = pagina_nl.locator(
        "#menun4"
    ).get_by_role("link")
    botao_confirmacao.click()

    verificar_panico()

    botao_confirmar = pagina_nl.get_by_role(
        "button",
        name="Confirmar a Operação",
    )
    botao_confirmar.wait_for(state="visible", timeout=10000)
    print('  [5/6] Confirmando a liquidação no SIGEF...', flush=True)
    botao_confirmar.click()

    mensagem_detalhada("Verificando mensagens de erro após a confirmação...")

    erro_nao_permitido = pagina_nl.get_by_text(
        "Não é permitido liquidar da"
    )

    try:
        erro_nao_permitido.wait_for(
            state="visible",
            timeout=5000,
        )

        print(
            "[AVISO] O SIGEF informou que a liquidação não é permitida "
            "neste fluxo ou já existe."
        )

        try:
            botao_voltar = pagina_nl.get_by_role(
                "button",
                name="Retornar à Tela Anterior",
            )
            botao_voltar.click()
        except Exception:
            pass

        salvar_nl(planilhas, registro, STATUS_PESQUISAR_NL,
        )

        raise ErroValidacaoSIGEF(
            "Linha marcada como PESQUISAR_NL para evitar "
            "uma possível liquidação duplicada."
        )

    except PlaywrightTimeoutError:
        pass

    aguardar_rede(pagina_nl)

    try:
        liquidacao = localizar_nl_gerada(pagina_nl)
    except Exception as exc:
        # A operação pode já ter sido confirmada; por segurança,
        # marca a linha para pesquisa e NÃO tenta novamente.
        salvar_nl(planilhas, registro, STATUS_PESQUISAR_NL,
        )
        raise ErroValidacaoSIGEF(
            "A operação foi confirmada, mas o número da NL não pôde "
            "ser capturado. A linha foi marcada como PESQUISAR_NL "
            "para evitar duplicidade."
        ) from exc

    print('  [6/6] Gravando a NL gerada na planilha...', flush=True)
    salvar_nl(planilhas, registro, liquidacao,
    )

    print(
        f"  [SALVO] {liquidacao} "
    )

    try:
        botao_limpar = pagina_nl.get_by_role(
            "link",
            name="Limpar a Tela",
        )
        botao_limpar.click()
    except Exception:
        pass



def executar_liquidacao(frame, guia, planilhas, ultima_linha: int):
    registros = [ler_registro(planilhas, linha)
                 for linha in range(2, ultima_linha + 1)
                 if not linha_vazia(planilhas["entrada"], linha)]
    existentes = sum(bool(r["liquidacao"]) and r["liquidacao"] != STATUS_PESQUISAR_NL
                     for r in registros)
    revisar = [r for r in registros if r["liquidacao"] == STATUS_PESQUISAR_NL
               or r["despesa_certificada"] == STATUS_PESQUISAR_CE]
    sem_ce = sum(not r["despesa_certificada"] and not r["liquidacao"] for r in registros)
    pendentes = [r for r in registros if r["despesa_certificada"]
                 and r["despesa_certificada"] != STATUS_PESQUISAR_CE and not r["liquidacao"]]
    print("\n" + "=" * 64)
    print("LIQUIDAÇÃO DE DESPESAS")
    print(f"Já liquidadas: {existentes} | A liquidar: {len(pendentes)} ")
    if sem_ce:
        print(f"Aguardando certificação: {sem_ce}")
    print("=" * 64)
    for r in revisar:
        documento = "NL" if r["liquidacao"] == STATUS_PESQUISAR_NL else "CE"
        print(f"[ATENÇÃO] Linha {r['linha']} | {r['nome']} | Localizar {documento} no SIGEF.")
    if not pendentes:
        print("Nenhuma nova liquidação a executar.")
        return
    pagina_nl = abrir_funcionalidade(frame, guia, "Liquidar Despesa Certificada")
    concluidas = 0
    try:
        for indice, registro in enumerate(pendentes, start=1):
            verificar_panico()
            print(f"\n[{indice}/{len(pendentes)}] Linha {registro['linha']} | {registro['nome']}")
            registro = validar_e_formatar_registro(registro)
            exibir_dados_registro(registro, "liquidacao")
            liquidar_linha(pagina_nl, planilhas, registro)
            concluidas += 1
    finally:
        try:
            pagina_nl.close()
        except Exception:
            pass
        print(f"\n[RESUMO] Liquidações concluídas nesta etapa: {concluidas} de {len(pendentes)}")


def mostrar_inicio(procedimento):
    if procedimento == "pp":
        nome_procedimento = "PREPARAR PAGAMENTOS"
        tarefas = (
            "• Consultar as Notas de Liquidação pendentes.\n"
            "• Conferir o favorecido e o domicílio bancário.\n"
            "• Gerar e salvar as Preparações de Pagamento na planilha."
        )
    else:
        nome_procedimento = "CERTIFICAR E LIQUIDAR"
        tarefas = (
            "• Gerar e salvar as Despesas Certificadas pendentes.\n"
            "• Liquidar as despesas que possuem CE válida.\n"
            "• Salvar os números das CEs e NLs na planilha."
        )

    instrucoes = (
        "ANTES DE INICIAR:\n"
        "• Salve e feche a planilha no Excel.\n"
        "• Entre no SIGEF pelo Chrome aberto pelo inicializador.\n"
        "• Deixe somente o menu principal do SIGEF aberto.\n"
        "• Execute apenas um robô por vez.\n\n"
        "DURANTE A EXECUÇÃO:\n"
        "• Não abra a planilha no Excel.\n"
        "• Não mova nem renomeie os arquivos do robô.\n"
        "• Não troque de aba nem clique nas janelas usadas pelo robô.\n"
        "• Não feche o Chrome ou as janelas abertas pelo SIGEF.\n"
        "• Interaja somente quando o robô solicitar uma seleção manual.\n"
        "• Para interromper, pressione ESC uma vez e aguarde."
    )
    return pyautogui.confirm(
        text=instrucoes,
        title=(
            "PREPARAR PAGAMENTO"
            if procedimento == "pp"
            else "CERTIFICAR E LIQUIDAR"
        ),
        buttons=["Iniciar", "Cancelar"],
    ) == "Iniciar"


def executar():
    planilhas = None
    hotkey_registrada = False

    try:
        if not mostrar_inicio("ce_nl"):
            print("[CANCELADO] Início cancelado. Nenhuma operação realizada.")
            return

        planilhas = abrir_planilha()
        ultima_linha = prevalidar_planilha(planilhas)

        keyboard.add_hotkey(
            TECLA_DE_PANICO,
            parar_execucao,
        )
        hotkey_registrada = True

        print(
            f"Robô iniciado. Pressione '{TECLA_DE_PANICO}' "
            "a qualquer momento para abortar com segurança."
        )


        verificar_panico()

        with sync_playwright() as p:
            print(
                "Tentando conectar ao Chrome na porta de depuração: "
                f"{CHROME_DEBUG_URL}"
            )

            browser = p.chromium.connect_over_cdp(
                CHROME_DEBUG_URL
            )

            print("Conexão estabelecida com sucesso!")

            if not browser.contexts:
                raise ErroAutomacao(
                    "Nenhum contexto do Chrome foi encontrado."
                )

            contexto = browser.contexts[0]
            guia = encontrar_guia_sigef(contexto)

            print(
                "Assumindo o controle da página: "
                f"'{guia.title()}'"
            )

            # Evita deixar o ano do caminho fixo.
            # Procura qualquer iframe de controle de acesso do SIGEF.
            frame = guia.frame_locator(
                'iframe[src*="SEGControleAcesso?p=1"]'
            )

            # Garante que o iframe realmente existe.
            frame.get_by_placeholder(
                "Pesquisar funcionalidades do sistema..."
            ).wait_for(
                state="visible",
                timeout=20000,
            )

            executar_certificacao(
                frame,
                guia,
                planilhas,
                ultima_linha,
            )

            verificar_panico()

            executar_liquidacao(
                frame,
                guia,
                planilhas,
                ultima_linha,
            )

        print(
            "\nScript finalizado. "
            "A janela de depuração permanece aberta."
        )

        pyautogui.alert(
            text="Encerrei por aqui.",
            title="Fim",
            button="OK",
        )

    except KeyboardInterrupt:
        print(
            "\n[ENCERRADO] Tecla ESC acionada. "
            "A planilha será fechada com segurança."
        )

        pyautogui.alert(
            text="Tecla ESC acionada. Automação encerrada.",
            title="Tecla de Pânico",
            button="OK",
        )

    except Exception as erro:
        print("\n" + "=" * 60)
        print("[ERRO] A automação foi interrompida por segurança.")
        print(f"Detalhes: {erro}")
        print("=" * 60)

        try:
            pyautogui.alert(
                text=(
                    "A automação foi interrompida para evitar "
                    f"registro incorreto.\n\n{erro}"
                ),
                title="Erro",
                button="OK",
            )
        except Exception:
            pass

        # Mantém traceback no terminal para diagnóstico.
        raise

    finally:
        if planilhas and planilhas.get("book"):
            try:
                planilhas["book"].close()
            except Exception:
                pass

        if hotkey_registrada:
            try:
                keyboard.remove_hotkey(TECLA_DE_PANICO)
            except Exception:
                pass

        print("[FINALIZAÇÃO] Recursos liberados com segurança.")


if __name__ == "__main__":
    executar()
